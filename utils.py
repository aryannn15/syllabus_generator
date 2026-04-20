import re
from functools import lru_cache

import fitz
import numpy as np
from sentence_transformers import SentenceTransformer

from config import BLOOM_RULES, EMBEDDING_MODEL, QA_MARKERS, TOPIC_PROTOTYPES, TOPICS
import os
from docx import Document

NOISE_PATTERNS = [
    "school of computer science",
    "continuous assessment test",
    "winter semester",
    "course name",
    "class number",
    "faculty name",
    "exam duration",
    "maximum marks",
    "programme name",
    "all questions carry equal marks",
    "answer all questions",
    "q.no.",
    "question max marks",
    "end of key",
]

QUESTION_CUE_PATTERN = (
    r"(?:^|\n)\s*(?:\d+\s*)?(?:[a-z]\)\s*)?"
    r"(?:find|explain|define|prove|design|convert|analyze|arrange|list|write|solve|implement|discuss|draw|construct|provide|perform|devise|develop|consider|compute|estimate|illustrate|mention|show|narrate|evaluate)\b"
)
ANSWER_MARKER_PATTERN = r"(?:^|\n)\s*(?:answer|ans\.?|solution|sol\.?)\s*[:.-]?\s*"
INLINE_ANSWER_MARKER_PATTERN = re.compile(r"^\s*(?:answer|ans\.?|solution|sol\.?)\s*[:.-]?\s*(.*)$", re.IGNORECASE)
SUBQUESTION_START_PATTERN = re.compile(r"^\s*(?:(\d+)\s*)?([a-h])\)\s*(.*)$", re.IGNORECASE)
QUESTION_START_PATTERN = re.compile(r"^\s*(\d+)\s*[.)]\s*(.*)$", re.IGNORECASE)

DIRECT_TOPIC_RULES = [
    ("Recurrence Relations", ["master theorem", "master method", "back substitution", "backward substitution", "t(n)", "recursive factorial"]),
    ("Asymptotic Notation", ["big o", "big omega", "theta", "o(", "o(g(n))", "n0", "c*n", "omega(", "θ(", "prove the following given function"]),
    ("Algorithm Correctness", ["proof of correctness", "prove the correctness", "loop invariant", "correctness of algorithm"]),
    ("Algorithm Basics", ["properties of an algorithm", "define algorithm", "algorithm development", "stages of algorithm"]),
    ("Linear Search", ["linear search"]),
    ("Linked List", ["linked list", "singly linked list", "doubly linked list", "circular single linked list", "create two lists", "display both the lists", "insertend", "insertinbetween", "deletenode"]),
    ("Binary Trees", ["expression tree", "generic tree", "general tree", "terminal nodes", "terminal node", "non-terminal", "internal nodes", "internal node", "strictly binary tree", "complete binary tree"]),
    ("Insertion Sort", ["insertion sort"]),
    ("Bubble Sort", ["bubble sort", "bubblesort"]),
    ("Quick Sort", ["quicksort", "quick sort"]),
    ("Merge Sort", ["merge sort"]),
    ("Binary Search", ["binary search"]),
    ("Counting Sort", ["counting sort"]),
    ("AVL Rotations", ["avl"]),
    ("Stack", [" stack", "stack ", "infix", "postfix", "push", "pop", "reverse polish"]),
    ("Queue", ["circular queue", " queue", "queue ", "enqueue", "dequeue"]),
    ("Hash Functions", ["hash function"]),
    ("Collision Resolution", ["collision"]),
    ("Separate Chaining", ["separate chaining", "chaining"]),
    ("Open Addressing", ["open addressing"]),
    ("Dijkstra", ["dijkstra", "single source shortest path", "shortest path"]),
    ("Minimum Spanning Tree", ["minimum spanning tree", "mst"]),
    ("Kruskal", ["kruskal"]),
    ("Prim", ["prim", "prim's algorithm"]),
    ("BFS", [" bfs", "breadth first"]),
    ("DFS", [" dfs", "depth first"]),
    ("Priority Queue", ["priority queue"]),
    ("Heap", ["heap"]),
    ("BST", ["bst", "binary search tree"]),
    ("Binary Trees", ["binary tree"]),
    ("Tree Traversals", ["traversal", "inorder", "preorder", "postorder", "level order"]),
    ("2-3 Trees", ["2-3 tree", "2 3 tree"]),
    ("Polynomial Manipulation", ["polynomial"]),
    ("Arrays", ["one dimensional array", " array "]),
    ("Time Complexity", ["growth rate", "worst-case", "worst case", "best case", "average case", "time complexity"]),
]

TOPIC_PRIORITY = [
    "Asymptotic Notation",
    "Recurrence Relations",
    "Arrays",
    "Stack",
    "Queue",
    "Linked List",
    "Polynomial Manipulation",
    "Algorithm Basics",
    "Algorithm Correctness",
    "Linear Search",
    "Binary Search",
    "Insertion Sort",
    "Bubble Sort",
    "Merge Sort",
    "Quick Sort",
    "Counting Sort",
    "Complexity Analysis (sorting-specific)",
    "Binary Trees",
    "BST",
    "Tree Traversals",
    "2-3 Trees",
    "Heap",
    "Priority Queue",
    "AVL Rotations",
    "BFS",
    "DFS",
    "Dijkstra",
    "Minimum Spanning Tree",
    "Prim",
    "Kruskal",
    "Hash Functions",
    "Collision Resolution",
    "Separate Chaining",
    "Open Addressing",
    "Time Complexity",
    "UNKNOWN",
]


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()

    # 🔹 PDF
    if ext == ".pdf":
        doc = fitz.open(file_path)
        text_chunks = []
        for page in doc:
            text_chunks.append(page.get_text())
        return "\n".join(text_chunks)

    # 🔹 DOCX (clean and reliable)
    elif ext == ".docx":
        return extract_docx_text(file_path)

    # 🔹 DOC (reject properly)
    elif ext == ".doc":
        print(f"❌ .doc not supported: {file_path}")
        print("👉 Convert to .docx and retry")
        return ""

    else:
        print(f"⚠️ Unsupported file: {file_path}")
        return ""


def extract_docx_text(file_path):
    doc = Document(file_path)
    text_chunks = []

    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            text_chunks.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_chunks = []
            seen_cells = set()
            for cell in row.cells:
                cell_lines = [
                    normalize_text(paragraph.text)
                    for paragraph in cell.paragraphs
                    if normalize_text(paragraph.text)
                ]
                cell_text = "\n".join(cell_lines)
                if cell_text and cell_text not in seen_cells:
                    row_chunks.append(cell_text)
                    seen_cells.add(cell_text)
            if row_chunks:
                text_chunks.append("\n".join(row_chunks))

    return "\n".join(text_chunks)

@lru_cache(maxsize=1)
def get_embedding_model():
    return SentenceTransformer(EMBEDDING_MODEL, local_files_only=True)


@lru_cache(maxsize=1)
def get_topic_prototype_embeddings():
    model = get_embedding_model()
    topic_names = []
    prototype_texts = []

    for topic, prototypes in TOPIC_PROTOTYPES.items():
        for prototype in prototypes:
            topic_names.append(topic)
            prototype_texts.append(prototype)

    embeddings = model.encode(prototype_texts, normalize_embeddings=True)
    return topic_names, embeddings


def normalize_text(text):
    return re.sub(r"\s+", " ", text).strip()

def normalize_math_text(text):
    text = text.lower()
    text = text.replace("nlogn", "n log n")
    text = text.replace("logn", "log n")
    text = re.sub(r"n(\d)", r"n^\1", text)  # n2 → n^2
    return text


def prepare_text(text):
    text = text.replace("\r", "\n")
    text = re.sub(r"\s+(\d+\s+[a-z]\))", r"\n\1", text, flags=re.IGNORECASE)
    text = re.sub(r"\n{2,}", "\n", text)
    return text


def is_noise_line(line):
    line_lower = line.lower()
    if len(line) < 2:
        return True
    if any(pattern in line_lower for pattern in NOISE_PATTERNS):
        return True
    if line in {"G2", "A1"}:
        return True
    if re.fullmatch(r"[0-9\s()./-]+", line):
        return True
    return False


def clean_text(text):
    text = prepare_text(text)
    lines = text.split("\n")
    filtered = []

    for raw_line in lines:
        line = normalize_text(raw_line)
        line_lower = line.lower()

        if len(line) < 8:
            continue
        if is_noise_line(line):
            continue
        if re.fullmatch(r".*\b\d+\s*marks?\b.*", line_lower) and len(line) < 25:
            continue

        filtered.append(line)

    cleaned = "\n".join(filtered)
    first_question = re.search(QUESTION_CUE_PATTERN, cleaned, flags=re.IGNORECASE)
    if first_question:
        cleaned = cleaned[first_question.start() :].lstrip()

    return cleaned


def split_questions(text):
    pattern = r"(?:^|\n)\s*(?:q(?:uestion)?\s*)?\d+\s*(?:[.)]\s+|[a-z]\)\s+)"
    parts = re.split(pattern, text, flags=re.IGNORECASE)
    questions = [normalize_text(part) for part in parts if len(normalize_text(part)) > 30]

    if questions:
        return questions

    return [normalize_text(block) for block in text.split("\n") if len(normalize_text(block)) > 30]


def is_marks_only_line(line):
    line_lower = line.lower()
    return bool(re.fullmatch(r"[()\d\s]*marks?[.)\s]*", line_lower))


def build_question_blocks(text):
    blocks = []
    current = None

    for raw_line in text.split("\n"):
        line = normalize_text(raw_line)
        if not line or line == "G2" or is_marks_only_line(line):
            continue

        subquestion_match = SUBQUESTION_START_PATTERN.match(line)
        if subquestion_match:
            if current:
                blocks.append(normalize_text(current))
            current = subquestion_match.group(3).strip()
            continue

        question_match = QUESTION_START_PATTERN.match(line)
        if question_match:
            if current:
                blocks.append(normalize_text(current))
            current = question_match.group(2).strip()
            continue

        if current:
            current = f"{current} {line}".strip()
        elif re.search(QUESTION_CUE_PATTERN, line, flags=re.IGNORECASE):
            current = line

    if current:
        blocks.append(normalize_text(current))

    return [block for block in blocks if len(block) > 15]


def detect_source_type(text):
    text_lower = text.lower()
    marker_hits = len(re.findall(ANSWER_MARKER_PATTERN, text_lower, flags=re.IGNORECASE))

    if marker_hits >= 2 or "answer key" in text_lower or "key slot" in text_lower or "end of key" in text_lower:
        return "question_answer"

    return "question_only"


def is_valid_question(text):
    text_lower = text.lower()

    if len(text_lower) < 16:
        return False
    if text_lower.startswith(":"):
        return False
    if "step 1" in text_lower:
        return False
    if "algorithm cq_" in text_lower:
        return False
    if "successive insertion of" in text_lower and "queue" not in text_lower and len(text_lower) < 70:
        return False
    if "successive deletion" in text_lower and "queue" not in text_lower and len(text_lower) < 70:
        return False
    if "faculty name" in text_lower or "exam duration" in text_lower:
        return False
    if text_lower.startswith("+") or text_lower.startswith("="):
        return False
    if text_lower.startswith("ans:") or text_lower.startswith("answer:") or text_lower.startswith("solution:"):
        return False
    if "falls under case" in text_lower or "applying big o" in text_lower:
        return False
    if re.fullmatch(r"[a-z ,.-]+", text_lower) and text_lower.count(",") >= 2:
        return False

    alpha_chars = sum(character.isalpha() for character in text)
    if alpha_chars < 8:
        return False

    return True


def is_answer_only_line(line):
    line_lower = line.lower()
    if INLINE_ANSWER_MARKER_PATTERN.match(line):
        return True
    if line_lower.startswith("applying big o"):
        return True
    if line_lower.startswith("falls under case"):
        return True
    if line_lower.startswith("substitute "):
        return True
    if line_lower.startswith("code count"):
        return True
    if line_lower.startswith("total"):
        return True
    if line_lower.startswith("let us assume"):
        return True
    if "under flow" in line_lower or "overflow" in line_lower:
        return True
    return False


def is_likely_answer_fragment(text):
    text_lower = text.lower().strip()
    if text_lower.startswith(("holds =", "applying ", "falls under case", "substitute ", "therefore ", "hence ")):
        return True
    if re.fullmatch(r"[0-9a-z\s().=+/*<>≤≥θ𝜃𝑛𝑓𝑔-]+", text_lower) and any(symbol in text_lower for symbol in ("=", "θ", "<=", ">=", "≤", "≥")):
        return not re.search(QUESTION_CUE_PATTERN, text, flags=re.IGNORECASE)

    return False


def should_attach_to_question(line):
    line_lower = line.lower()
    if is_marks_only_line(line):
        return False
    if re.match(r"^\d+\.\s+(successive|insertion|deletion)", line_lower):
        return True
    if re.match(r"^[ivx]+\)", line_lower):
        return True
    if "t(n)" in line_lower or "f(n)" in line_lower or "o(" in line_lower or "θ(" in line_lower:
        return True
    if "log" in line_lower or "n2" in line_lower or "n3" in line_lower:
        return True
    if "for(" in line_lower or "sum" in line_lower or "a[k]" in line_lower:
        return True
    if len(line) < 10 and any(char.isdigit() for char in line):
        return False
    return False


def compact_answer_text(answer_lines):
    kept = []
    for line in answer_lines:
        normalized = normalize_text(line)
        if not normalized:
            continue
        if len(normalized) < 12:
            continue
        if re.fullmatch(r"[0-9\s().=+/*<>-]+", normalized):
            continue
        if normalized.lower().startswith(("ans", "answer", "solution")):
            continue
        kept.append(normalized)
        if len(kept) == 3:
            break

    return normalize_text(" ".join(kept))


def build_question_answer_entries(text):
    lines = [normalize_text(line) for line in prepare_text(text).split("\n")]
    entries = []
    current_question = []
    current_answer = []
    in_answer = False

    def flush_entry():
        question_text = normalize_text(" ".join(current_question))
        answer_text = compact_answer_text(current_answer)
        if is_likely_answer_fragment(question_text) and re.search(QUESTION_CUE_PATTERN, answer_text, flags=re.IGNORECASE):
            question_text, answer_text = answer_text, question_text
        question_lower = question_text.lower()
        if question_lower.startswith("successive insertion") or question_lower.startswith("successive deletion"):
            return
        if question_lower.startswith("insertion of the element") and not answer_text:
            return
        if is_valid_question(question_text):
            entries.append(
                {
                    "question_text": question_text,
                    "answer_text": answer_text,
                }
            )

    for line in lines:
        if not line or is_noise_line(line):
            continue
        if is_marks_only_line(line):
            continue

        subquestion_match = SUBQUESTION_START_PATTERN.match(line)
        question_match = QUESTION_START_PATTERN.match(line)
        if subquestion_match or question_match:
            if current_question and "instructions in sequence" in " ".join(current_question).lower():
                current_question.append(subquestion_match.group(3).strip() if subquestion_match else question_match.group(2).strip())
                continue
            if current_question:
                flush_entry()
            current_question = [subquestion_match.group(3).strip() if subquestion_match else question_match.group(2).strip()]
            current_answer = []
            in_answer = False
            continue

        answer_match = INLINE_ANSWER_MARKER_PATTERN.match(line)
        if answer_match and current_question:
            in_answer = True
            answer_fragment = normalize_text(answer_match.group(1))
            if answer_fragment:
                current_answer.append(answer_fragment)
            continue

        if not current_question:
            continue

        if not in_answer:
            if should_attach_to_question(line):
                current_question.append(line)
            elif is_answer_only_line(line):
                in_answer = True
                current_answer.append(line)
            elif re.search(QUESTION_CUE_PATTERN, line, flags=re.IGNORECASE):
                flush_entry()
                current_question = [line]
                current_answer = []
                in_answer = False
            else:
                current_answer.append(line)
                in_answer = True
        else:
            if is_answer_only_line(line) or len(line) > 12:
                current_answer.append(line)

    if current_question:
        flush_entry()

    return entries


def parse_entries(text, source_type):
    if source_type == "question_answer":
        return build_question_answer_entries(text)

    entries = []
    cleaned_text = clean_text(text)
    questions = build_question_blocks(cleaned_text)
    if not questions:
        questions = split_questions(cleaned_text)

    for question in questions:
        part = normalize_text(question)
        if is_valid_question(part):
            entries.append(
                {
                    "question_text": part,
                    "answer_text": "",
                }
            )

    return entries


def detect_direct_topic(text):
    text_lower = f" {text.lower()} "

    if ("n0" in text_lower or "𝑛0" in text_lower) and ("o(g" in text_lower or "𝑜(𝑔" in text_lower or "𝑂(𝑔" in text_lower or "f(n)" in text_lower or "𝑓(𝑛)" in text_lower):
        return "Asymptotic Notation"
    if ("find c" in text_lower or "find c and" in text_lower) and ("n0" in text_lower or "𝑛0" in text_lower):
        return "Asymptotic Notation"

    for topic, keywords in DIRECT_TOPIC_RULES:
        for keyword in keywords:
            if keyword_matches(text_lower, keyword.lower()):
                return topic

    return None


def is_weak_embedding_guess(primary_topic, scores, rule_scores, embedding_scores):
    if primary_topic in rule_scores:
        return False

    primary_score = float(scores.get(primary_topic, 0))
    primary_embedding = float(embedding_scores.get(primary_topic, 0))
    if primary_score < 1.75:
        return True
    if primary_embedding < 0.5 and primary_score < 2.2:
        return True

    return False


def keyword_matches(text, keyword):
    if any(char in keyword for char in "()[]=+*/<>θ"):
        return keyword in text

    escaped = re.escape(keyword.strip())
    if " " in keyword.strip() or "-" in keyword.strip():
        pattern = rf"(?<![a-z0-9]){escaped}(?![a-z0-9])"
    else:
        pattern = rf"\b{escaped}\b"

    return bool(re.search(pattern, text))


def compute_topic_scores(question_text, answer_text=""):
    question_lower = f" {question_text.lower()} "
    answer_lower = f" {answer_text.lower()} "
    scores = {}

    for topic, keywords in TOPICS.items():
        score = 0
        for keyword in keywords:
            keyword_lower = keyword.lower()
            if keyword_matches(question_lower, keyword_lower):
                score += 2
            if answer_text and keyword_matches(answer_lower, keyword_lower):
                score += 1
        if score > 0:
            scores[topic] = score

    return scores


def compute_embedding_scores(question_text, answer_text=""):
    try:
        model = get_embedding_model()
        topic_names, prototype_embeddings = get_topic_prototype_embeddings()
        if answer_text:
            combined_text = f"{question_text} || context: {answer_text}"
        else:
            combined_text = question_text
        query_embedding = model.encode([combined_text], normalize_embeddings=True)[0]
        similarities = np.dot(prototype_embeddings, query_embedding)

        topic_scores = {}
        for topic, similarity in zip(topic_names, similarities):
            score = float(similarity)
            if topic not in topic_scores or score > topic_scores[topic]:
                topic_scores[topic] = score

        return topic_scores
    except Exception:
        return {}


def blend_scores(rule_scores, embedding_scores):
    blended = {}
    all_topics = set(rule_scores) | set(embedding_scores)

    for topic in all_topics:
        rule_component = float(rule_scores.get(topic, 0))
        embedding_component = float(embedding_scores.get(topic, 0))
        blended_score = (rule_component*1.2) + (embedding_component * 3.5)
        if blended_score > 0:
            blended[topic] = round(blended_score, 4)

    return blended


def choose_secondary_topics(scores, primary_topic):
    if not scores or primary_topic == "UNKNOWN":
        return []

    primary_score = float(scores.get(primary_topic, 0))
    secondary = []

    for topic, score in sorted(scores.items(), key=lambda item: (-item[1], item[0])):
        if topic == primary_topic:
            continue
        if score >= max(1.0, primary_score - 0.75):
            secondary.append(topic)
        elif score >= 2.5 and primary_score >= 4:
            secondary.append(topic)
        if len(secondary) == 2:
            break

    return secondary


def topic_priority(topic):
    try:
        return TOPIC_PRIORITY.index(topic)
    except ValueError:
        return len(TOPIC_PRIORITY)


def choose_primary_topic(scores):
    max_score = max(scores.values())
    candidates = [topic for topic, score in scores.items() if score == max_score]
    candidates.sort(key=topic_priority)
    return candidates[0]


def compact_topic_scores(blended_scores, rule_scores, embedding_scores):
    compact = {}
    kept = 0
    for topic, total_score in sorted(blended_scores.items(), key=lambda item: (-item[1], item[0])):
        rule_score = float(rule_scores.get(topic, 0))
        embedding_score = float(embedding_scores.get(topic, 0))
        if kept >= 4 and rule_score == 0 and total_score < 2.0:
            continue
        if rule_score == 0 and total_score < 1.5:
            continue
        compact[topic] = {
            "score": round(total_score, 2),
            "rule": round(rule_score, 2),
            "embedding": round(embedding_score, 3),
        }
        kept += 1
    return compact


def compute_confidence(question_text, answer_text, primary_topic, scores, used_direct_match, embedding_scores):
    if primary_topic == "UNKNOWN":
        return 0.2

    signal = float(scores.get(primary_topic, 0))
    embedding_signal = float(embedding_scores.get(primary_topic, 0))
    question_len = len(question_text.split())
    answer_bonus = 0.1 if answer_text else 0.0
    direct_bonus = 0.2 if used_direct_match else 0.0
    signal_bonus = min(0.35, signal * 0.05)
    embedding_bonus = min(0.25, max(0.0, embedding_signal) * 0.35)
    length_bonus = 0.1 if question_len >= 8 else 0.0
    confidence = 0.25 + direct_bonus + signal_bonus + embedding_bonus + length_bonus + answer_bonus

    return round(min(confidence, 0.98), 2)


def map_topic(question_text, answer_text=""):
    question_text = normalize_math_text(question_text)
    answer_text = normalize_math_text(answer_text)
    rule_scores = compute_topic_scores(question_text, answer_text)
    embedding_scores = compute_embedding_scores(question_text, answer_text)
    direct_topic = detect_direct_topic(question_text)

    if direct_topic:
        rule_scores[direct_topic] = max(rule_scores.get(direct_topic, 0), 4)

    scores = blend_scores(rule_scores, embedding_scores)
    if not scores:
        return "UNKNOWN", [], {}, 0.2

    primary = choose_primary_topic(scores)
    if is_weak_embedding_guess(primary, scores, rule_scores, embedding_scores):
        compact_scores = compact_topic_scores(scores, rule_scores, embedding_scores)
        return "UNKNOWN", [], compact_scores, 0.2

    secondary = choose_secondary_topics(scores, primary)
    compact_scores = compact_topic_scores(scores, rule_scores, embedding_scores)
    confidence = compute_confidence(question_text, answer_text, primary, scores, bool(direct_topic), embedding_scores)

    return primary, secondary, compact_scores, confidence


def detect_bloom(text):
    text_lower = text.lower()

    for level, keywords in BLOOM_RULES.items():
        for keyword in keywords:
            if keyword in text_lower:
                return level

    return "Understand"
